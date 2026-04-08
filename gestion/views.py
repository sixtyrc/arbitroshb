from rest_framework import viewsets, permissions, generics, status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response
from rest_framework.views import APIView
from django.http import HttpResponse
from django.conf import settings
from .models import CustomUser, Arbitro, Partido, Disponibilidad, Designacion, PushSubscription
from .serializers import (
    UserSerializer, ArbitroSerializer, ArbitroProfileUpdateSerializer,
    PartidoSerializer, DisponibilidadSerializer, DesignacionSerializer
)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

# ── ViewSets Estándar ─────────────────────────────────────────────────────────

class UserViewSet(viewsets.ModelViewSet):
    queryset = CustomUser.objects.all()
    serializer_class = UserSerializer
    permission_classes = [permissions.IsAdminUser]

class ArbitroViewSet(viewsets.ModelViewSet):
    queryset = Arbitro.objects.all()
    serializer_class = ArbitroSerializer

    def get_permissions(self):
        if self.action in ['list', 'retrieve']:
            return [permissions.IsAuthenticated()]
        return [permissions.IsAdminUser()]

class PartidoViewSet(viewsets.ModelViewSet):
    queryset = Partido.objects.all()
    serializer_class = PartidoSerializer

    def get_permissions(self):
        if self.action in ['list', 'retrieve']:
            return [permissions.IsAuthenticated()]
        return [permissions.IsAdminUser()]

class DisponibilidadViewSet(viewsets.ModelViewSet):
    queryset = Disponibilidad.objects.all()
    serializer_class = DisponibilidadSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'ADMIN':
            return Disponibilidad.objects.all()
        return Disponibilidad.objects.filter(arbitro__user=user)

class DesignacionViewSet(viewsets.ModelViewSet):
    queryset = Designacion.objects.all()
    serializer_class = DesignacionSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'ADMIN':
            return Designacion.objects.all()
        return Designacion.objects.filter(arbitro__user=user)


# ── Perfil propio del Árbitro (GET y PATCH) ──────────────────────────────────

class MiPerfilView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        try:
            arbitro = request.user.perfil_arbitro
        except Arbitro.DoesNotExist:
            return Response({'detail': 'Perfil no encontrado.'}, status=404)
        serializer = ArbitroSerializer(arbitro)
        return Response(serializer.data)

    def patch(self, request):
        try:
            arbitro = request.user.perfil_arbitro
        except Arbitro.DoesNotExist:
            return Response({'detail': 'Perfil no encontrado.'}, status=404)
        serializer = ArbitroProfileUpdateSerializer(arbitro, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


# ── Liquidaciones Excel (solo ADMIN) ─────────────────────────────────────────

class LiquidacionesExcelView(APIView):
    permission_classes = [permissions.IsAdminUser]

    def get(self, request):
        # Obtener parámetros de filtro opcionales
        mes = request.query_params.get('mes')    # formato: YYYY-MM
        aceptadas_only = request.query_params.get('aceptadas', 'true').lower() == 'true'

        designaciones = Designacion.objects.select_related(
            'arbitro__user', 'partido__categoria__grupo'
        ).all()

        if aceptadas_only:
            designaciones = designaciones.filter(status='ACCEPTED')
        if mes:
            try:
                year, month = mes.split('-')
                designaciones = designaciones.filter(
                    partido__date_time__year=int(year),
                    partido__date_time__month=int(month)
                )
            except ValueError:
                pass

        # ── Crear workbook ────────────────────────────────────────────────────
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Liquidaciones"

        # Estilos
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill("solid", fgColor="1e3a5f")
        center = Alignment(horizontal='center', vertical='center')
        accent_fill = PatternFill("solid", fgColor="10b981")

        # Título principal
        ws.merge_cells('A1:H1')
        titulo = ws['A1']
        titulo.value = f"COLEGIO DE ÁRBITROS DE HANDBALL DEL CHACO — Liquidación {mes or 'Completa'}"
        titulo.font = Font(bold=True, size=13, color="FFFFFF")
        titulo.fill = PatternFill("solid", fgColor="0f172a")
        titulo.alignment = center
        ws.row_dimensions[1].height = 30

        # Subtítulo
        ws.merge_cells('A2:H2')
        sub = ws['A2']
        sub.value = f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        sub.font = Font(italic=True, color="888888")
        sub.alignment = center

        # Headers
        headers = ['Árbitro', 'Email', 'CBU / Alias', 'Partido', 'Categoría', 'Fecha', 'Rol', 'Honorario ($)']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        # Datos
        total = 0
        for row_idx, des in enumerate(designaciones, start=4):
            arbitro = des.arbitro
            partido = des.partido
            monto = float(des.monto_honorario)
            total += monto

            row_data = [
                arbitro.user.get_full_name(),
                arbitro.user.email,
                arbitro.cbu_alias or '⚠ Sin CBU',
                partido.title,
                partido.categoria.nombre if partido.categoria else 'S/C',
                partido.date_time.strftime('%d/%m/%Y %H:%M'),
                des.get_rol_asignado_display(),
                monto,
            ]
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.alignment = Alignment(vertical='center')
                # Alternar color de fila
                if row_idx % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="f0f9ff")

        # Fila Total
        total_row = len(list(designaciones)) + 4
        ws.cell(row=total_row, column=7, value="TOTAL:").font = Font(bold=True)
        total_cell = ws.cell(row=total_row, column=8, value=total)
        total_cell.font = Font(bold=True, color="FFFFFF")
        total_cell.fill = accent_fill
        total_cell.alignment = center

        # Anchos de columna
        col_widths = [25, 28, 22, 30, 15, 18, 18, 15]
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

        # Respuesta HTTP
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"liquidacion_{mes or 'completa'}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response


# ── Web Push: Llave pública y suscripción ─────────────────────────────────────

class VapidPublicKeyView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        return Response({'public_key': settings.VAPID_PUBLIC_KEY})


class PushSubscribeView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        data = request.data
        endpoint = data.get('endpoint')
        p256dh = data.get('keys', {}).get('p256dh', '')
        auth = data.get('keys', {}).get('auth', '')

        if not endpoint:
            return Response({'error': 'Endpoint requerido.'}, status=status.HTTP_400_BAD_REQUEST)

        PushSubscription.objects.update_or_create(
            user=request.user,
            defaults={'endpoint': endpoint, 'p256dh': p256dh, 'auth': auth}
        )
        return Response({'detail': 'Suscripción guardada.'}, status=status.HTTP_201_CREATED)
